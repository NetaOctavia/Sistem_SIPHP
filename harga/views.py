import csv
from datetime import timedelta
import json
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.core.paginator import Paginator
from django.core.serializers.json import DjangoJSONEncoder
from django.db.models import Avg
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from accounts.models import Pasar
from berita.models import Berita
from komoditas.models import Komoditas
from .models import HargaKomoditas

KATEGORI_EKSPOR = "Komoditi Ekspor"


def beranda(request):
    selected_pasar_id = request.GET.get("pasar_id", "")
    daftar_pasar = Pasar.objects.all().order_by("nama_pasar")

    # Ambil entri harga terbaru
    harga_qs = HargaKomoditas.objects.all()
    if selected_pasar_id:
        harga_qs = harga_qs.filter(pasar_id=selected_pasar_id)

    latest_entry = harga_qs.order_by("-tanggal").first()
    today = latest_entry.tanggal if latest_entry else timezone.now().date()
    yesterday = today - timedelta(days=1)

    avg_harga_query = HargaKomoditas.objects.filter(tanggal=today)
    if selected_pasar_id:
        avg_harga_query = avg_harga_query.filter(pasar_id=selected_pasar_id)

    avg_harga = avg_harga_query.aggregate(Avg("harga"))["harga__avg"] or 0
    berita_list = Berita.objects.all().order_by("-created_at")[:3]

    harga_today = HargaKomoditas.objects.filter(tanggal=today).select_related(
        "komoditas", "pasar"
    )
    if selected_pasar_id:
        harga_today = harga_today.filter(pasar_id=selected_pasar_id)

    data_harga_formatted = []

    # Batch fetch semua harga kemarin (1 query)
    harga_yesterday_qs = HargaKomoditas.objects.filter(tanggal=yesterday).select_related(
        "komoditas", "pasar"
    )
    if selected_pasar_id:
        harga_yesterday_qs = harga_yesterday_qs.filter(pasar_id=selected_pasar_id)
    harga_yesterday_map = {}
    for hy in harga_yesterday_qs:
        harga_yesterday_map[(hy.komoditas_id, hy.pasar_id)] = hy.harga

    for h in harga_today:
        key = (h.komoditas_id, h.pasar_id)
        harga_kemarin_val = harga_yesterday_map.get(key, 0)
        has_yesterday = key in harga_yesterday_map
        perubahan = h.harga - harga_kemarin_val if has_yesterday else 0

        data_harga_formatted.append({
            "nama_item": (
                f"{h.komoditas.nama} ({h.pasar.nama_pasar})"
                if h.pasar
                else h.komoditas.nama
            ),
            "satuan": h.komoditas.satuan,
            "harga_kemarin": harga_kemarin_val,
            "harga_hari_ini": h.harga,
            "perubahan": perubahan,
            "perubahan_abs": abs(perubahan),
        })

    # LOGIKA GRAFIK TREN HARGA BERDASARKAN PERIODE
    selected_periode = request.GET.get("periode", "7")
    if selected_periode == "14":
        start_date = today - timedelta(days=13)
        periode_label = "14 Hari Terakhir"
    elif selected_periode == "30":
        start_date = today - timedelta(days=29)
        periode_label = "30 Hari Terakhir"
    elif selected_periode == "bulan_ini":
        start_date = today.replace(day=1)
        periode_label = f"Bulan Ini ({today.strftime('%B %Y')})"
    else:
        selected_periode = "7"
        start_date = today - timedelta(days=6)
        periode_label = "7 Hari Terakhir"

    delta_days = (today - start_date).days + 1
    dates_list = [(start_date + timedelta(days=i)) for i in range(delta_days)]
    dates_label = [d.strftime("%d %b") for d in dates_list]

    top_names = [
        "Beras Premium", "Cabe Rawit Merah", "Bawang Merah",
        "Daging Sapi", "Daging Ayam Boiler", "Telur Ayam Boiler"
    ]
    daftar_komoditas_plot = Komoditas.objects.filter(nama__in=top_names).order_by("nama")
    if not daftar_komoditas_plot.exists():
        daftar_komoditas_plot = Komoditas.objects.all()[:6]

    colors = [
        {"border": "#10B981", "bg": "rgba(16, 185, 129, 0.08)"},
        {"border": "#EF4444", "bg": "rgba(239, 68, 68, 0.08)"},
        {"border": "#3B82F6", "bg": "rgba(59, 130, 246, 0.08)"},
        {"border": "#F59E0B", "bg": "rgba(245, 158, 11, 0.08)"},
        {"border": "#8B5CF6", "bg": "rgba(139, 92, 246, 0.08)"},
        {"border": "#EC4899", "bg": "rgba(236, 72, 153, 0.08)"},
    ]

    # Batch fetch harga rentang periode (1 query)
    chart_harga_qs = HargaKomoditas.objects.filter(
        komoditas__in=daftar_komoditas_plot,
        tanggal__range=(start_date, today)
    )
    if selected_pasar_id:
        chart_harga_qs = chart_harga_qs.filter(pasar_id=selected_pasar_id)
    chart_agg = chart_harga_qs.values(
        'komoditas_id', 'tanggal'
    ).annotate(avg_harga=Avg('harga'))
    chart_harga_map = {}
    for item in chart_agg:
        chart_harga_map[(item['komoditas_id'], item['tanggal'])] = float(item['avg_harga'])

    chart_datasets = []
    for idx, kom in enumerate(daftar_komoditas_plot):
        data_harga_per_hari = [
            chart_harga_map.get((kom.id, d), 0) for d in dates_list
        ]

        color_scheme = colors[idx % len(colors)]
        chart_datasets.append({
            "label": kom.nama,
            "data": data_harga_per_hari,
            "borderColor": color_scheme["border"],
            "backgroundColor": color_scheme["bg"],
            "fill": False,
            "tension": 0.35,
            "pointRadius": 3,
            "pointHoverRadius": 5,
        })

    context = {
        "avg_harga": round(avg_harga),
        "tanggal_terbaru": today,
        "tanggal_kemarin": yesterday,
        "data_harga": data_harga_formatted,
        "berita_list": berita_list,
        "daftar_pasar": daftar_pasar,
        "selected_pasar_id": selected_pasar_id,
        "selected_periode": selected_periode,
        "periode_label": periode_label,
        "komoditas_naik": sum(
            1 for item in data_harga_formatted if item["perubahan"] > 0
        ),
        "komoditas_turun": sum(
            1 for item in data_harga_formatted if item["perubahan"] < 0
        ),
        "chart_labels": json.dumps(dates_label, cls=DjangoJSONEncoder),
        "chart_datasets": json.dumps(chart_datasets, cls=DjangoJSONEncoder),
    }
    return render(request, "dashboard/beranda.html", context)


@staff_member_required(login_url='login')
def dashboard_index(request):
    """
    Overview dashboard admin SIPHP ( /dashboard/ )
    """
    total_pasar = Pasar.objects.count()
    total_komoditas = Komoditas.objects.count()
    today = timezone.now().date()
    total_input_today = HargaKomoditas.objects.filter(tanggal=today).count()

    status_pasar_list = []
    for p in Pasar.objects.all().order_by("nama_pasar"):
        c = HargaKomoditas.objects.filter(pasar=p, tanggal=today).count()
        status_pasar_list.append({
            "pasar": p,
            "sudah_input": c > 0,
            "jumlah_input": c,
        })

    context = {
        "total_pasar": total_pasar,
        "total_komoditas": total_komoditas,
        "total_input_today": total_input_today,
        "status_pasar_list": status_pasar_list,
    }
    return render(request, "admin_dashboard/index.html", context)


@staff_member_required(login_url='login')
def harga_komoditas(request):
    user_pasar = None
    if hasattr(request.user, "profil_admin") and request.user.profil_admin and request.user.profil_admin.pasar:
        user_pasar = request.user.profil_admin.pasar

    if request.method == "POST":
        # 1. Form Tambah Pasar Baru (Khusus Admin Utama)
        if "tambah_pasar" in request.POST:
            if user_pasar:
                messages.error(request, "Anda tidak memiliki akses untuk menambah pasar.")
                return redirect("harga_komoditas")
            nama_pasar_baru = request.POST.get("nama_pasar_baru")
            if nama_pasar_baru:
                pasar_obj, _ = Pasar.objects.get_or_create(
                    nama_pasar=nama_pasar_baru.strip()
                )
                messages.success(
                    request, f'Pasar "{pasar_obj.nama_pasar}" berhasil ditambahkan!'
                )
                return redirect(f"/dashboard/harga/?pasar_id={pasar_obj.id}")
            return redirect("harga_komoditas")

        # 2. Form Tambah Komoditas Baru
        if "tambah_komoditas" in request.POST:
            if user_pasar:
                messages.error(request, "Hanya Admin Utama yang dapat menambah master komoditas.")
                return redirect("harga_komoditas")
            nama_komoditas_baru = request.POST.get("nama_komoditas")
            satuan_baru = request.POST.get("satuan", "kg")
            kategori_baru = (request.POST.get("kategori") or "").strip()
            harga_awal = request.POST.get("harga_awal", "")

            if nama_komoditas_baru:
                kom_obj, created = Komoditas.objects.get_or_create(
                    nama=nama_komoditas_baru.strip(),
                    defaults={
                        "satuan": satuan_baru.strip(),
                        "kategori": kategori_baru or None,
                    },
                )

                if harga_awal:
                    val_clean = "".join(filter(str.isdigit, harga_awal))
                    if val_clean:
                        harga_val = int(val_clean)
                        pasar_target = user_pasar or Pasar.objects.first()
                        if pasar_target:
                            HargaKomoditas.objects.update_or_create(
                                komoditas=kom_obj,
                                pasar=pasar_target,
                                tanggal=timezone.now().date(),
                                defaults={
                                    "harga": harga_val,
                                    "diinput_oleh": request.user,
                                },
                            )
                messages.success(
                    request, f'Komoditas "{kom_obj.nama}" berhasil ditambahkan!'
                )
            return redirect("harga_komoditas")

        # 3. Form Batch Update Harga
        pasar_id = request.POST.get("pasar_id")
        tanggal_str = request.POST.get("tanggal")

        # Jika user terikat ke pasar tertentu, paksakan pasar tersebut
        if user_pasar:
            pasar_obj = user_pasar
        else:
            pasar_obj = get_object_or_404(Pasar, id=pasar_id)

        try:
            tanggal = (
                timezone.datetime.strptime(tanggal_str, "%Y-%m-%d").date()
                if tanggal_str
                else timezone.now().date()
            )
        except ValueError:
            tanggal = timezone.now().date()

        count_updated = 0

        for key, val in request.POST.items():
            if key.startswith("harga_") and val.strip():
                try:
                    komoditas_id = int(key.replace("harga_", ""))
                    val_clean = "".join(filter(str.isdigit, val))

                    if val_clean:
                        harga_val = int(val_clean)
                        komoditas_obj = Komoditas.objects.get(id=komoditas_id)

                        HargaKomoditas.objects.update_or_create(
                            komoditas=komoditas_obj,
                            pasar=pasar_obj,
                            tanggal=tanggal,
                            defaults={
                                "harga": harga_val,
                                "diinput_oleh": request.user,
                            },
                        )
                        count_updated += 1
                except (ValueError, Komoditas.DoesNotExist):
                    continue

        if count_updated > 0:
            messages.success(
                request,
                f"Berhasil memperbarui harga {count_updated} komoditas di {pasar_obj.nama_pasar}!",
            )
        else:
            messages.warning(
                request, "Tidak ada harga komoditas yang dimasukkan atau diubah."
            )

        return redirect(f"/dashboard/harga/?pasar_id={pasar_obj.id}&tanggal={tanggal}")

    # --- TAMPILAN GET REQUEST ---
    daftar_pasar = Pasar.objects.all().order_by("nama_pasar")
    selected_pasar_id = request.GET.get("pasar_id")
    selected_tanggal = request.GET.get(
        "tanggal", timezone.now().date().strftime("%Y-%m-%d")
    )
    selected_kategori = request.GET.get("kategori", "").strip()
    search_q = request.GET.get("q", "").strip()

    pasar_aktif = None
    if user_pasar:
        pasar_aktif = user_pasar
    elif selected_pasar_id:
        pasar_aktif = Pasar.objects.filter(id=selected_pasar_id).first()
    elif daftar_pasar.exists():
        pasar_aktif = daftar_pasar.first()

    daftar_komoditas = Komoditas.objects.all().order_by("nama")
    if selected_kategori:
        daftar_komoditas = daftar_komoditas.filter(kategori=selected_kategori)
    if search_q:
        daftar_komoditas = daftar_komoditas.filter(nama__icontains=search_q)

    list_komoditas_dengan_harga = []

    harga_existing_map = {}
    harga_terakhir_map = {}

    if pasar_aktif:
        for h in HargaKomoditas.objects.filter(
            pasar=pasar_aktif, tanggal=selected_tanggal
        ):
            harga_existing_map[h.komoditas_id] = int(h.harga)

        for h in HargaKomoditas.objects.filter(
            pasar=pasar_aktif
        ).order_by("tanggal", "id"):
            harga_terakhir_map[h.komoditas_id] = int(h.harga)

    for kom in daftar_komoditas:
        list_komoditas_dengan_harga.append({
            "id": kom.id,
            "nama": kom.nama,
            "kategori": kom.kategori or "",
            "satuan": kom.satuan or "kg",
            "harga_existing": harga_existing_map.get(kom.id),
            "harga_terakhir": harga_terakhir_map.get(kom.id, 0),
        })

    # Pagination 10 komoditas per halaman
    paginator = Paginator(list_komoditas_dengan_harga, 10)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    # Pilihan kategori: kategori yang sudah dipakai + kategori baru yang diminta
    daftar_kategori = sorted(
        set(
            Komoditas.objects.exclude(kategori__isnull=True)
            .exclude(kategori="")
            .values_list("kategori", flat=True)
        )
    )
    if KATEGORI_EKSPOR not in daftar_kategori:
        daftar_kategori.insert(0, KATEGORI_EKSPOR)

    return render(
        request,
        "auth_custom/kelola_harga.html",
        {
            "komoditas_list": page_obj,
            "page_obj": page_obj,
            "daftar_pasar": daftar_pasar,
            "pasar_aktif": pasar_aktif,
            "tanggal_input": selected_tanggal,
            "today_date": timezone.now().date().strftime("%Y-%m-%d"),
            "daftar_kategori": daftar_kategori,
            "selected_kategori": selected_kategori,
            "search_q": search_q,
        },
    )


@staff_member_required(login_url='login')
def hapus_pasar(request, pasar_id):
    if request.method != "POST":
        return redirect("harga_komoditas")
    pasar = get_object_or_404(Pasar, id=pasar_id)
    nama_pasar = pasar.nama_pasar
    pasar.delete()
    messages.success(request, f'Pasar "{nama_pasar}" berhasil dihapus.')
    return redirect("harga_komoditas")


@staff_member_required(login_url='login')
def export_harga_csv(request):
    selected_pasar_id = request.GET.get("pasar_id")
    selected_tanggal = request.GET.get("tanggal")

    qs = HargaKomoditas.objects.select_related("komoditas", "pasar", "diinput_oleh").all()
    if selected_pasar_id:
        qs = qs.filter(pasar_id=selected_pasar_id)
    if selected_tanggal:
        qs = qs.filter(tanggal=selected_tanggal)

    qs = qs.order_by("-tanggal", "pasar__nama_pasar", "komoditas__nama")

    response = HttpResponse(content_type="text/csv; charset=utf-8")
    filename = f"rekap_harga_pangan_{selected_tanggal or timezone.now().date().strftime('%Y-%m-%d')}.csv"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    # Write BOM for Excel UTF-8 compatibility
    response.write("\ufeff".encode("utf8"))
    writer = csv.writer(response)
    writer.writerow(["No", "Tanggal", "Nama Pasar", "Komoditas", "Satuan", "Harga (Rp)", "Petugas Input"])

    for idx, item in enumerate(qs, start=1):
        writer.writerow([
            idx,
            item.tanggal.strftime("%Y-%m-%d"),
            item.pasar.nama_pasar if item.pasar else "-",
            item.komoditas.nama if item.komoditas else "-",
            item.komoditas.satuan if item.komoditas else "kg",
            int(item.harga),
            item.diinput_oleh.username if item.diinput_oleh else "-",
        ])

    return response


def _clean_price_value(val):
    if val is None:
        return None
    val_str = str(val).strip()
    val_str = val_str.replace("Rp", "").replace("rp", "").replace(" ", "")
    if "." in val_str and "," in val_str:
        val_str = val_str.split(",")[0].replace(".", "")
    elif "." in val_str:
        parts = val_str.split(".")
        if len(parts) == 2 and len(parts[1]) <= 2:
            val_str = parts[0]
        else:
            val_str = val_str.replace(".", "")
    elif "," in val_str:
        parts = val_str.split(",")
        if len(parts) == 2 and len(parts[1]) <= 2:
            val_str = parts[0]
        else:
            val_str = val_str.replace(",", "")
    try:
        num = int(float(val_str))
        return num if num >= 0 else None
    except (ValueError, TypeError):
        return None


def _parse_date_flexible(val, fallback_date):
    import datetime
    if val is None or val == "":
        return fallback_date
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.date() if isinstance(val, datetime.datetime) else val
    val_str = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.datetime.strptime(val_str, fmt).date()
        except ValueError:
            pass
    return fallback_date


@staff_member_required(login_url='login')
def import_harga_file(request):
    """
    Fitur Bulk Import Data Harga Komoditas dari file CSV atau Excel (.xlsx).
    Mendukung logika:
    - Tambah data baru jika belum ada.
    - Update harga jika komoditas & pasar & tanggal sudah ada tapi harganya berbeda.
    - Lewati (skip) jika harga sama agar tidak duplikat.
    - Proteksi hak akses pasar untuk Admin Pasar.
    - Deteksi baris header secara fleksibel.
    """
    if request.method != "POST":
        return redirect("harga_komoditas")

    uploaded_file = request.FILES.get("file_import")
    if not uploaded_file:
        messages.error(request, "Silakan pilih file CSV atau Excel (.xlsx) untuk diimpor.")
        return redirect("harga_komoditas")

    filename = uploaded_file.name.lower()
    if not (filename.endswith(".csv") or filename.endswith(".xlsx") or filename.endswith(".xls")):
        messages.error(request, "Format file tidak didukung. Harap unggah file berformat .csv atau .xlsx.")
        return redirect("harga_komoditas")

    # Hak akses pasar penugasan pengguna
    user_pasar = None
    if hasattr(request.user, "profil_admin") and request.user.profil_admin and request.user.profil_admin.pasar:
        user_pasar = request.user.profil_admin.pasar

    target_pasar_id = request.POST.get("target_pasar_id")
    target_pasar_default = None
    if target_pasar_id:
        target_pasar_default = Pasar.objects.filter(id=target_pasar_id).first()

    if user_pasar:
        target_pasar_default = user_pasar

    default_tanggal_str = request.POST.get("default_tanggal")
    fallback_tanggal = _parse_date_flexible(default_tanggal_str, timezone.now().date())

    rows_data = []

    try:
        if filename.endswith(".csv"):
            import io
            content = uploaded_file.read()
            decoded = None
            for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
                try:
                    decoded = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue

            if decoded is None:
                messages.error(request, "Gagal membaca encoding file CSV. Pastikan file tersimpan dalam format UTF-8.")
                return redirect("harga_komoditas")

            lines = [l for l in decoded.splitlines() if l.strip()]
            delimiter = ","
            if lines:
                for sample_line in lines[:3]:
                    count_semicolon = sample_line.count(";")
                    count_comma = sample_line.count(",")
                    count_tab = sample_line.count("\t")
                    if count_semicolon > count_comma and count_semicolon > count_tab:
                        delimiter = ";"
                        break
                    elif count_tab > count_comma and count_tab > count_semicolon:
                        delimiter = "\t"
                        break
                    elif count_comma > 0:
                        delimiter = ","
                        break

            csv_reader = csv.reader(io.StringIO(decoded), delimiter=delimiter)
            for r in csv_reader:
                if any(c.strip() for c in r):
                    rows_data.append([str(c).strip() for c in r])

        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(values_only=True):
                if row and any(c is not None and str(c).strip() != "" for c in row):
                    clean_row = [str(c).strip() if c is not None else "" for c in row]
                    rows_data.append(clean_row)

    except Exception as e:
        messages.error(request, f"Terjadi kesalahan saat memproses file: {str(e)}")
        return redirect("harga_komoditas")

    if not rows_data:
        messages.error(request, "File yang Anda unggah kosong atau tidak memiliki data yang dapat dibaca.")
        return redirect("harga_komoditas")

    header_indices = {
        "tanggal": None,
        "pasar": None,
        "komoditas": None,
        "kategori": None,
        "satuan": None,
        "harga": None,
    }

    header_row_index = -1

    for r_idx, row in enumerate(rows_data[:6]):
        row_lower = [c.lower() for c in row]
        matches = 0
        temp_indices = {"tanggal": None, "pasar": None, "komoditas": None, "kategori": None, "satuan": None, "harga": None}

        for idx, col in enumerate(row_lower):
            if any(k in col for k in ["tanggal", "tgl", "date", "waktu"]):
                temp_indices["tanggal"] = idx
                matches += 1
            elif any(k in col for k in ["nama pasar", "pasar", "market", "lokasi"]):
                temp_indices["pasar"] = idx
                matches += 1
            elif any(k in col for k in ["komoditas", "nama komoditas", "nama barang", "commodity"]):
                temp_indices["komoditas"] = idx
                matches += 1
            elif any(k in col for k in ["kategori", "category", "jenis"]):
                temp_indices["kategori"] = idx
                matches += 1
            elif any(k in col for k in ["satuan", "unit", "uom"]):
                temp_indices["satuan"] = idx
                matches += 1
            elif any(k in col for k in ["harga", "price", "nominal", "nilai"]):
                temp_indices["harga"] = idx
                matches += 1

        if (temp_indices["komoditas"] is not None or temp_indices["harga"] is not None) and matches >= 2:
            header_indices = temp_indices
            header_row_index = r_idx
            break

    if header_row_index != -1:
        data_rows = rows_data[header_row_index + 1:]
        has_header = True
    else:
        has_header = False
        data_rows = rows_data

    semua_pasar = list(Pasar.objects.all())
    pasar_by_name = {p.nama_pasar.lower().strip(): p for p in semua_pasar}

    created_count = 0
    updated_count = 0
    skipped_count = 0
    error_list = []

    for row_idx, row in enumerate(data_rows, start=header_row_index + 2 if has_header else 1):
        if not row or not any(row):
            continue

        first_cell_lower = str(row[0]).lower().strip()
        if first_cell_lower.startswith("petunjuk") or first_cell_lower.startswith("catatan") or first_cell_lower.startswith("template"):
            continue

        if has_header:
            val_tanggal = row[header_indices["tanggal"]] if header_indices["tanggal"] is not None and header_indices["tanggal"] < len(row) else ""
            val_pasar = row[header_indices["pasar"]] if header_indices["pasar"] is not None and header_indices["pasar"] < len(row) else ""
            val_komoditas = row[header_indices["komoditas"]] if header_indices["komoditas"] is not None and header_indices["komoditas"] < len(row) else ""
            val_kategori = row[header_indices["kategori"]] if header_indices["kategori"] is not None and header_indices["kategori"] < len(row) else ""
            val_satuan = row[header_indices["satuan"]] if header_indices["satuan"] is not None and header_indices["satuan"] < len(row) else "kg"
            val_harga = row[header_indices["harga"]] if header_indices["harga"] is not None and header_indices["harga"] < len(row) else ""
        else:
            if len(row) >= 6:
                val_tanggal, val_pasar, val_komoditas, val_kategori, val_satuan, val_harga = row[0], row[1], row[2], row[3], row[4], row[5]
            elif len(row) >= 4:
                val_tanggal, val_pasar, val_komoditas, val_kategori, val_satuan, val_harga = row[0], "", row[1], "", row[2], row[3]
            elif len(row) == 2:
                val_tanggal, val_pasar, val_komoditas, val_kategori, val_satuan, val_harga = "", "", row[0], "", "kg", row[1]
            else:
                continue

        nama_komoditas = str(val_komoditas).strip()
        if not nama_komoditas or nama_komoditas.lower() in ("komoditas", "nama komoditas"):
            continue

        harga_parsed = _clean_price_value(val_harga)
        if harga_parsed is None:
            error_list.append(f"Baris {row_idx} ({nama_komoditas}): Harga tidak valid ('{val_harga}').")
            continue

        row_tanggal = _parse_date_flexible(val_tanggal, fallback_tanggal)

        row_pasar = None
        str_pasar = str(val_pasar).strip()

        if user_pasar:
            row_pasar = user_pasar
            if str_pasar and str_pasar.lower() != user_pasar.nama_pasar.lower().strip():
                error_list.append(f"Baris {row_idx} ({nama_komoditas}): Anda hanya diizinkan mengimpor data untuk pasar '{user_pasar.nama_pasar}'.")
                continue
        else:
            if str_pasar:
                row_pasar = pasar_by_name.get(str_pasar.lower())
                if not row_pasar:
                    for p_name, p_obj in pasar_by_name.items():
                        if str_pasar.lower() in p_name or p_name in str_pasar.lower():
                            row_pasar = p_obj
                            break
                if not row_pasar:
                    row_pasar = Pasar.objects.create(nama_pasar=str_pasar)
                    semua_pasar.append(row_pasar)
                    pasar_by_name[row_pasar.nama_pasar.lower().strip()] = row_pasar
            else:
                row_pasar = target_pasar_default

        if not row_pasar:
            error_list.append(f"Baris {row_idx} ({nama_komoditas}): Pasar tidak ditentukan. Silakan isi kolom pasar atau pilih pasar di formulir.")
            continue

        satuan_clean = str(val_satuan).strip().lower() or "kg"
        kategori_clean = str(val_kategori).strip() or None

        komoditas_obj = Komoditas.objects.filter(
            Q(nama__iexact=nama_komoditas, pasar=row_pasar) | Q(nama__iexact=nama_komoditas, pasar__isnull=True)
        ).first()

        if not komoditas_obj:
            komoditas_obj = Komoditas.objects.create(
                nama=nama_komoditas,
                pasar=row_pasar,
                kategori=kategori_clean,
                satuan=satuan_clean,
            )

        harga_exist = HargaKomoditas.objects.filter(
            komoditas=komoditas_obj,
            pasar=row_pasar,
            tanggal=row_tanggal
        ).first()

        if harga_exist:
            if int(harga_exist.harga) != int(harga_parsed):
                harga_exist.harga = harga_parsed
                harga_exist.diinput_oleh = request.user
                harga_exist.save()
                updated_count += 1
            else:
                skipped_count += 1
        else:
            HargaKomoditas.objects.create(
                komoditas=komoditas_obj,
                pasar=row_pasar,
                tanggal=row_tanggal,
                harga=harga_parsed,
                diinput_oleh=request.user
            )
            created_count += 1

    msg_parts = []
    if created_count > 0:
        msg_parts.append(f"✨ {created_count} harga baru ditambahkan")
    if updated_count > 0:
        msg_parts.append(f"🔄 {updated_count} harga diperbarui")
    if skipped_count > 0:
        msg_parts.append(f"⏭️ {skipped_count} data dilewati (harga tetap sama)")

    if msg_parts:
        summary_msg = "Proses import selesai: " + ", ".join(msg_parts) + "."
        messages.success(request, summary_msg)
    elif not error_list:
        messages.info(request, "Seluruh data dalam file sudah sesuai dan tidak ada perubahan harga.")

    if error_list:
        max_show = 5
        error_sample = "; ".join(error_list[:max_show])
        if len(error_list) > max_show:
            error_sample += f" ...dan {len(error_list) - max_show} baris lainnya."
        messages.warning(request, f"Terdapat {len(error_list)} baris yang gagal diproses: {error_sample}")

    redirect_url = request.META.get("HTTP_REFERER") or "harga_komoditas"
    return redirect(redirect_url)


@staff_member_required(login_url='login')
def download_template_import(request, format_type="csv"):
    """
    Menyediakan template unduhan untuk format CSV atau Excel (.xlsx).
    """
    headers = ["Tanggal", "Nama Pasar", "Komoditas", "Kategori", "Satuan", "Harga"]
    contoh_data = [
        ["2026-09-22", "Pasar Pujasera Subang", "Beras Medium", "Beras & Padi", "kg", 13500],
        ["2026-09-22", "Pasar Pujasera Subang", "Bawang Merah", "Cabai & Bumbu", "kg", 35000],
        ["2026-09-22", "Pasar Pujasera Subang", "Cabe Merah Keriting", "Cabai & Bumbu", "kg", 42000],
        ["2026-09-22", "Pasar Pujasera Subang", "Daging Sapi Paha Belakang", "Peternakan & Daging", "kg", 130000],
        ["2026-09-22", "Pasar Pujasera Subang", "Minyak Goreng Kemasan", "Minyak & Lemak", "Liter", 18000],
        ["2026-09-22", "Pasar Pujasera Subang", "Telur Ayam Boiler", "Peternakan & Daging", "kg", 29000],
        ["2026-09-22", "Pasar Pujasera Subang", "Gula Pasir", "Sembako & Olahan", "kg", 17500],
    ]

    if format_type.lower() == "csv":
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="template_import_harga_siphp.csv"'
        response.write("\ufeff".encode("utf8"))

        writer = csv.writer(response, delimiter=";")
        writer.writerow(headers)
        for row in contoh_data:
            writer.writerow(row)
        return response

    elif format_type.lower() in ("excel", "xlsx"):
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Template Harga SIPHP"
        ws.views.sheetView[0].showGridLines = True

        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = "SIPHP SUBANG — TEMPLATE IMPORT HARGA KOMODITAS PASAR"
        title_cell.font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="007A4D", end_color="007A4D", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 32

        ws.merge_cells("A2:F2")
        info_cell = ws["A2"]
        info_cell.value = "💡 Petunjuk: Isi data komoditas mulai baris ke-4. Jangan ubah nama kolom header pada baris ke-3."
        info_cell.font = Font(name="Segoe UI", size=9, italic=True, color="1E3A2F")
        info_cell.fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
        info_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[2].height = 22

        ws.append(headers)
        ws.row_dimensions[3].height = 26

        header_fill = PatternFill(start_color="009661", end_color="009661", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        border_thin = Side(style="thin", color="D0D5DD")
        border_all = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_all

        zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        data_font = Font(name="Segoe UI", size=10, color="1F2937")

        for r_offset, row in enumerate(contoh_data, start=4):
            ws.append(row)
            ws.row_dimensions[r_offset].height = 20
            row_fill = zebra_fill if (r_offset % 2 == 0) else white_fill

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_offset, column=col_idx)
                cell.font = data_font
                cell.border = border_all
                cell.fill = row_fill

                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in (2, 3, 4):
                    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                elif col_idx == 5:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 6:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '#,##0'

        column_widths = {"A": 16, "B": 26, "C": 28, "D": 20, "E": 12, "F": 18}
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        import io
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="template_import_harga_siphp.xlsx"'
        return response

    return redirect("harga_komoditas")



def api_harga_komoditas(request):
    selected_pasar_id = request.GET.get("pasar_id", "")
    selected_tanggal = request.GET.get("tanggal", "")

    qs = HargaKomoditas.objects.select_related("komoditas", "pasar").all()
    if selected_pasar_id:
        qs = qs.filter(pasar_id=selected_pasar_id)
    if selected_tanggal:
        qs = qs.filter(tanggal=selected_tanggal)

    data = []
    for item in qs.order_by("-tanggal", "komoditas__nama")[:100]:
        data.append({
            "id": item.id,
            "komoditas": item.komoditas.nama,
            "satuan": item.komoditas.satuan,
            "pasar": item.pasar.nama_pasar,
            "harga": int(item.harga),
            "tanggal": item.tanggal.strftime("%Y-%m-%d"),
        })

    return JsonResponse({"status": "success", "count": len(data), "data": data}, safe=False)
